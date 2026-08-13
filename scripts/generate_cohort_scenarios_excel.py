"""Generate cohort-scenarios.xlsx — simple, plain-language scenarios.

Only two sheets: Dashboard and Cohorts page. Kept intentionally short so it's
easy to scan and fill in during a review call.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = "cohort-scenarios.xlsx"

HEADERS = ["#", "Who is the user?", "Situation", "What happens today", "What it should do"]

DASHBOARD_ROWS = [
    (1, "New user (not enrolled anywhere)", "No cohort is open for registration right now",
     "Card says 'No cohorts available right now' with a Browse cohorts button."),
    (2, "New user", "A cohort is upcoming and open for registration",
     "Card shows that cohort with an Upcoming tag and a Register button."),
    (3, "New user", "No upcoming cohort is open, but an older one still accepts registrations (dates already passed)",
     "Card shows that cohort with an Enrollment open tag and a View details button (not Register)."),
    (4, "Enrolled user", "Their cohort is running right now",
     "Card labeled Active cohort, with a progress ring and a View details button."),
    (5, "Enrolled user", "Their cohort hasn't started yet",
     "Card labeled Your cohort with Upcoming + Enrolled tags, no progress ring, View details button."),
    (6, "Enrolled user", "Their cohort ended and they didn't finish everything",
     "Card labeled Your cohort with an Ended tag and a progress ring."),
    (7, "Enrolled user", "Their cohort ended and they finished 100%",
     "Card labeled Your cohort with a Completed tag and a full progress ring."),
    (8, "Any user (new or enrolled)", "There's a cohort open for registration that isn't their own",
     "A second, smaller card below their own always shows it — skipped only if it's the exact same cohort they're already in."),
]

COHORTS_PAGE_ROWS = [
    (1, "New user", "Cohort hasn't started yet, registration is open",
     "Shown under 'Upcoming Cohorts' with an Open tag and a Register button."),
    (2, "New user", "Cohort is running right now, registration is open",
     "Shown under 'Ongoing Cohorts' with an Open tag."),
    (3, "New user", "Cohort dates already ended, but registration is still open",
     "Shown under 'Open Enrollment' with an Open tag and a View details button (not Register, since the dates already passed)."),
    (4, "New user (never part of this cohort)", "Cohort dates ended and registration is closed",
     "Not shown at all — no point cluttering the page with a cohort this user was never part of and can't join."),
    (5, "Enrolled user", "Their cohort is running right now",
     "Shown under 'Ongoing Cohorts' with an Enrolled or Completed tag and an Open button."),
    (6, "Enrolled user", "Their cohort hasn't started yet",
     "Shown under 'Upcoming Cohorts' with an Enrolled tag and an Open button."),
    (7, "Enrolled user", "Their cohort already ended",
     "Shown under 'Past Cohorts' with an Enrolled or Completed tag and an Open button."),
    (8, "Admin", "Any past cohort, enrolled or not",
     "Always shown under 'Past Cohorts' — admins need full visibility, so they're exempt from rule #4 above. Admins also see draft/unpublished cohorts everywhere on this page."),
]


def style_header(ws, num_cols, row=1):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF")
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autosize(ws, max_width=55):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), max_width)


def write_sheet(ws, rows):
    for c, h in enumerate(HEADERS, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, len(HEADERS))

    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=len(HEADERS), value="")  # blank "should do"

    ws.freeze_panes = "A2"
    autosize(ws, max_width=50)
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 50
    ws.column_dimensions["E"].width = 35


def main():
    wb = Workbook()

    dashboard = wb.active
    dashboard.title = "Dashboard"
    write_sheet(dashboard, DASHBOARD_ROWS)

    cohorts_page = wb.create_sheet("Cohorts page")
    write_sheet(cohorts_page, COHORTS_PAGE_ROWS)

    wb.save(OUTPUT)
    print(f"Created {OUTPUT} with 2 sheets ({len(DASHBOARD_ROWS)} + {len(COHORTS_PAGE_ROWS)} rows)")


if __name__ == "__main__":
    main()
